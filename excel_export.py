"""
Экспорт отчётов в Excel.
"""

import os
from datetime import date
from typing import List, Dict, Optional

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


class ExcelExporter:
    def export(self, records: List[Dict], export_date: date) -> str:
        """Создаёт Excel файл и возвращает путь к нему."""
        if not OPENPYXL_AVAILABLE:
            return self._export_csv(records, export_date)
        return self._export_xlsx(records, export_date)

    def export_multi_chat(self, chat_data: Dict[int, Dict], export_date: date) -> str:
        """
        Создаёт один Excel файл с листами по каждому чату.
        chat_data: {chat_id: {"records": [...], "stats": {...}, "fin": {...}}}
        """
        if not OPENPYXL_AVAILABLE:
            return self._export_csv_multi_chat(chat_data, export_date)
        return self._export_xlsx_multi_chat(chat_data, export_date)

    def _export_xlsx_multi_chat(self, chat_data: Dict[int, Dict], export_date: date) -> str:
        wb = openpyxl.Workbook()

        # Удаляем дефолтный лист
        default_sheet = wb.active
        wb.remove(default_sheet)

        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill("solid", fgColor="1F4E79")
        header_align = Alignment(horizontal="center", vertical="center")
        total_font = Font(bold=True, size=11)
        total_fill = PatternFill("solid", fgColor="D6E4F0")
        thin = Side(style="thin", color="CCCCCC")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        headers = ["№", "Время", "Сотрудник", "Номер машины", "Фреон (г/мл)", "Сумма (руб.)", "Оплата", "Кому", "Банк"]
        col_widths = [5, 12, 18, 16, 14, 14, 10, 14, 10]

        for chat_id, data in chat_data.items():
            records = data.get("records", [])
            refills = data.get("refills", [])
            stats = data.get("stats", {})
            fin = data.get("fin")

            # Берем имя чата из данных, если нет — используем ID
            sheet_name = data.get("name", str(chat_id))
            # Excel не любит символы []*?/\ и длину > 31
            for char in ['[', ']', '*', '?', '/', '\\', ':']:
                sheet_name = sheet_name.replace(char, '')
            ws = wb.create_sheet(title=sheet_name[:31])

            for col_idx, (header, width) in enumerate(zip(headers, col_widths), 1):
                cell = ws.cell(row=1, column=col_idx, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_align
                cell.border = border
                ws.column_dimensions[get_column_letter(col_idx)].width = width
            ws.row_dimensions[1].height = 22

            for row_idx, record in enumerate(records, 2):
                time_str = record.get("created_at", "")
                if "T" in str(time_str):
                    time_str = time_str.split("T")[1][:5]
                elif " " in str(time_str):
                    time_str = str(time_str).split(" ")[1][:5]

                payment_label = record.get("payment_type", "")
                if payment_label == "безнал" and record.get("payment_name") and record.get("payment_bank"):
                    payment_label += f" ({record['payment_name']}, {record['payment_bank']})"
                row_data = [
                    row_idx - 1,
                    time_str,
                    record.get("username", ""),
                    record.get("car_number", ""),
                    record.get("freon_volume", 0),
                    record.get("money", 0),
                    payment_label,
                    record.get("payment_name", ""),
                    record.get("payment_bank", ""),
                ]
                alt_fill = PatternFill("solid", fgColor="EBF3FB") if row_idx % 2 == 0 else None
                for col_idx, value in enumerate(row_data, 1):
                    cell = ws.cell(row=row_idx, column=col_idx, value=value)
                    cell.border = border
                    cell.alignment = Alignment(horizontal="center" if col_idx in (1, 2, 4, 5, 6, 7, 8, 9) else "left")
                    if alt_fill:
                        cell.fill = alt_fill

            total_row = len(records) + 2
            total_freon = sum(r.get("freon_volume", 0) for r in records)
            total_money = sum(r.get("money", 0) for r in records)

            ws.cell(row=total_row, column=1, value="ИТОГО").font = total_font
            ws.cell(row=total_row, column=1).fill = total_fill
            ws.cell(row=total_row, column=3, value=f"{len(records)} машин").font = total_font
            ws.cell(row=total_row, column=3).fill = total_fill
            ws.cell(row=total_row, column=5, value=total_freon).font = total_font
            ws.cell(row=total_row, column=5).fill = total_fill
            ws.cell(row=total_row, column=6, value=total_money).font = total_font
            ws.cell(row=total_row, column=6).fill = total_fill
            for col_idx in range(1, 10):
                ws.cell(row=total_row, column=col_idx).border = border
                ws.cell(row=total_row, column=col_idx).alignment = Alignment(horizontal="center")

            if refills:
                refill_start = total_row + 2
                ws.cell(row=refill_start, column=1, value="Заправки аппарата:").font = total_font
                for i, refill in enumerate(refills):
                    row = refill_start + 1 + i
                    time_str = refill.get("created_at", "")
                    if "T" in str(time_str):
                        time_str = time_str.split("T")[1][:5]
                    elif " " in str(time_str):
                        time_str = str(time_str).split(" ")[1][:5]
                    ws.cell(row=row, column=2, value=time_str).font = total_font
                    ws.cell(row=row, column=3, value=refill.get("username", "")).font = total_font
                    ws.cell(row=row, column=5, value=f"{refill.get('refill_kg', 0)} кг").font = total_font

            if fin:
                fin_row = total_row + 2 + (len(refills) + 1 if refills else 0)
                ws.cell(row=fin_row, column=1, value="Финансы:").font = total_font
                ws.cell(row=fin_row, column=3, value=f"Общая: {fin.get('fr_total', 0)}").font = total_font
                ws.cell(row=fin_row + 1, column=3, value=f"Мои: {fin.get('fr_mine', 0)}").font = total_font
                ws.cell(row=fin_row + 2, column=3, value=f"Перевожу: {fin.get('fr_transfer', 0)}").font = total_font

            # Безнал разбивка по точкам
            beznal_rows = [r for r in records if r.get("payment_type") == "безнал"]
            if beznal_rows:
                bd_header = total_row + 2 + (len(refills) + 2 if refills else 0) + (4 if fin else 0)
                ws.cell(row=bd_header, column=1, value="Безнал (разбивка):").font = total_font
                bd_headers = ["Кому", "Банк", "Сумма"]
                for ci, h in enumerate(bd_headers, 1):
                    c = ws.cell(row=bd_header + 1, column=ci, value=h)
                    c.font = header_font
                    c.fill = header_fill
                    c.border = border
                by_name = {}
                for r in beznal_rows:
                    key = (r.get("payment_name", "?"), r.get("payment_bank", ""))
                    by_name[key] = by_name.get(key, 0) + r.get("money", 0)
                for i, ((name, bank), total) in enumerate(sorted(by_name.items()), bd_header + 2):
                    bank_label = bank if bank else "—"
                    for ci, v in enumerate([name.capitalize() if name else "?", bank_label, total], 1):
                        c = ws.cell(row=i, column=ci, value=v)
                        c.border = border
                        c.alignment = Alignment(horizontal="center")

        # Лист "Итого" — общая сводка по всем чатам
        summary_ws = wb.create_sheet(title="Итого", index=0)
        summary_headers = ["Чат", "Машин", "Фреон (г)", "Выручка (руб.)", "Заправка (кг)"]
        summary_widths = [12, 12, 14, 18, 14]
        for col_idx, (header, width) in enumerate(zip(summary_headers, summary_widths), 1):
            cell = summary_ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = border
            summary_ws.column_dimensions[get_column_letter(col_idx)].width = width
        summary_ws.row_dimensions[1].height = 22

        grand_total_cars = 0
        grand_total_freon = 0
        grand_total_money = 0
        grand_total_refill = 0

        for row_idx, (chat_id, data) in enumerate(chat_data.items(), 2):
            stats = data.get("stats", {})
            cars = stats.get("count", 0)
            freon = stats.get("total_freon", 0)
            money = stats.get("total_money", 0)
            refill = stats.get("total_refill", 0)
            grand_total_cars += cars
            grand_total_freon += freon
            grand_total_money += money
            grand_total_refill += refill

            alt_fill = PatternFill("solid", fgColor="EBF3FB") if row_idx % 2 == 0 else None
            chat_name = data.get("name", str(chat_id))
            for col_idx, value in enumerate([chat_name, cars, freon, money, refill], 1):
                cell = summary_ws.cell(row=row_idx, column=col_idx, value=value)
                cell.border = border
                cell.alignment = Alignment(horizontal="center")
                if alt_fill:
                    cell.fill = alt_fill

        total_row = len(chat_data) + 2
        summary_ws.cell(row=total_row, column=1, value="ВСЕГО").font = total_font
        summary_ws.cell(row=total_row, column=1).fill = total_fill
        summary_ws.cell(row=total_row, column=2, value=grand_total_cars).font = total_font
        summary_ws.cell(row=total_row, column=2).fill = total_fill
        summary_ws.cell(row=total_row, column=3, value=grand_total_freon).font = total_font
        summary_ws.cell(row=total_row, column=3).fill = total_fill
        summary_ws.cell(row=total_row, column=4, value=grand_total_money).font = total_font
        summary_ws.cell(row=total_row, column=4).fill = total_fill
        summary_ws.cell(row=total_row, column=5, value=grand_total_refill).font = total_font
        summary_ws.cell(row=total_row, column=5).fill = total_fill
        for col_idx in range(1, 6):
            summary_ws.cell(row=total_row, column=col_idx).border = border
            summary_ws.cell(row=total_row, column=col_idx).alignment = Alignment(horizontal="center")

        # Лист "Безнал" — сводка по безналу по всем точкам
        bd_ws = wb.create_sheet(title="Безнал")
        bd_headers = ["Чат", "Кому", "Банк", "Сумма"]
        bd_widths = [14, 14, 12, 14]
        for ci, (h, w) in enumerate(zip(bd_headers, bd_widths), 1):
            c = bd_ws.cell(row=1, column=ci, value=h)
            c.font = header_font
            c.fill = header_fill
            c.alignment = header_align
            c.border = border
            bd_ws.column_dimensions[get_column_letter(ci)].width = w
        bd_ws.row_dimensions[1].height = 22

        bd_row = 2
        grand_bd_total = 0
        for chat_id, data in chat_data.items():
            records = data.get("records", [])
            beznal_rows = [r for r in records if r.get("payment_type") == "безнал"]
            if not beznal_rows:
                continue
            by_name = {}
            for r in beznal_rows:
                key = (r.get("payment_name", "?"), r.get("payment_bank", ""))
                by_name[key] = by_name.get(key, 0) + r.get("money", 0)
            chat_name = data.get("name", str(chat_id))
            for (name, bank), total in sorted(by_name.items()):
                bank_label = bank if bank else "—"
                for ci, v in enumerate([chat_name, name.capitalize() if name else "?", bank_label, total], 1):
                    c = bd_ws.cell(row=bd_row, column=ci, value=v)
                    c.border = border
                    c.alignment = Alignment(horizontal="center")
                grand_bd_total += total
                bd_row += 1

        if bd_row > 2:
            bd_ws.cell(row=bd_row, column=1, value="ВСЕГО").font = total_font
            bd_ws.cell(row=bd_row, column=1).fill = total_fill
            bd_ws.cell(row=bd_row, column=4, value=grand_bd_total).font = total_font
            bd_ws.cell(row=bd_row, column=4).fill = total_fill
            for ci in range(1, 5):
                bd_ws.cell(row=bd_row, column=ci).border = border
                bd_ws.cell(row=bd_row, column=ci).alignment = Alignment(horizontal="center")
        else:
            bd_ws.cell(row=2, column=1, value="Безналичных платежей нет")

        filename = f"report_{export_date.strftime('%Y-%m-%d')}.xlsx"
        wb.save(filename)
        return filename

    def _export_csv_multi_chat(self, chat_data: Dict[int, Dict], export_date: date) -> str:
        filename = f"report_{export_date.strftime('%Y-%m-%d')}.csv"
        with open(filename, "w", encoding="utf-8-sig") as f:
            for chat_id, data in chat_data.items():
                f.write(f"\n=== Чат {chat_id} ===\n")
                f.write("№,Время,Сотрудник,Номер машины,Фреон,Сумма,Оплата,Кому,Банк\n")
                records = data.get("records", [])
                for i, r in enumerate(records, 1):
                    f.write(f"{i},{r.get('created_at','')},{r.get('username','')},{r.get('car_number','')},{r.get('freon_volume',0)},{r.get('money',0)},{r.get('payment_type','')},{r.get('payment_name','')},{r.get('payment_bank','')}\n")
                total_freon = sum(r.get("freon_volume", 0) for r in records)
                total_money = sum(r.get("money", 0) for r in records)
                f.write(f"ИТОГО,,,{len(records)} машин,{total_freon},{total_money},\n")
        return filename

    def _export_xlsx(self, records: List[Dict], export_date: date) -> str:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = export_date.strftime("%d.%m.%Y")

        # Стили
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill("solid", fgColor="1F4E79")
        header_align = Alignment(horizontal="center", vertical="center")

        total_font = Font(bold=True, size=11)
        total_fill = PatternFill("solid", fgColor="D6E4F0")

        thin = Side(style="thin", color="CCCCCC")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        # Заголовки
        headers = ["№", "Время", "Сотрудник", "Номер машины", "Фреон (г/мл)", "Сумма (руб.)", "Оплата", "Кому", "Банк"]
        col_widths = [5, 12, 18, 16, 14, 14, 10, 14, 10]

        for col_idx, (header, width) in enumerate(zip(headers, col_widths), 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = border
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        ws.row_dimensions[1].height = 22

        # Данные
        for row_idx, record in enumerate(records, 2):
            time_str = record.get("created_at", "")
            if "T" in str(time_str):
                time_str = time_str.split("T")[1][:5]
            elif " " in str(time_str):
                time_str = str(time_str).split(" ")[1][:5]

            payment_label = record.get("payment_type", "")
            if payment_label == "безнал" and record.get("payment_name") and record.get("payment_bank"):
                payment_label += f" ({record['payment_name']}, {record['payment_bank']})"
            row_data = [
                row_idx - 1,
                time_str,
                record.get("username", ""),
                record.get("car_number", ""),
                record.get("freon_volume", 0),
                record.get("money", 0),
                payment_label,
                record.get("payment_name", ""),
                record.get("payment_bank", ""),
            ]

            alt_fill = PatternFill("solid", fgColor="EBF3FB") if row_idx % 2 == 0 else None

            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.border = border
                cell.alignment = Alignment(horizontal="center" if col_idx in (1, 2, 4, 5, 6, 7, 8, 9) else "left")
                if alt_fill:
                    cell.fill = alt_fill

        # Итоговая строка
        total_row = len(records) + 2
        total_freon = sum(r.get("freon_volume", 0) for r in records)
        total_money = sum(r.get("money", 0) for r in records)

        ws.cell(row=total_row, column=1, value="ИТОГО").font = total_font
        ws.cell(row=total_row, column=1).fill = total_fill
        ws.cell(row=total_row, column=3, value=f"{len(records)} машин").font = total_font
        ws.cell(row=total_row, column=3).fill = total_fill
        ws.cell(row=total_row, column=5, value=total_freon).font = total_font
        ws.cell(row=total_row, column=5).fill = total_fill
        ws.cell(row=total_row, column=6, value=total_money).font = total_font
        ws.cell(row=total_row, column=6).fill = total_fill

        for col_idx in range(1, 10):
            ws.cell(row=total_row, column=col_idx).border = border
            ws.cell(row=total_row, column=col_idx).alignment = Alignment(horizontal="center")

        # Сохраняем
        filename = f"report_{export_date.strftime('%Y-%m-%d')}.xlsx"
        wb.save(filename)
        return filename

    def _export_csv(self, records: List[Dict], export_date: date) -> str:
        """Резервный вариант — CSV."""
        filename = f"report_{export_date.strftime('%Y-%m-%d')}.csv"
        with open(filename, "w", encoding="utf-8-sig") as f:
            f.write("№,Время,Сотрудник,Номер машины,Фреон,Сумма,Оплата,Кому,Банк\n")
            for i, r in enumerate(records, 1):
                f.write(f"{i},{r.get('created_at','')},{r.get('username','')},{r.get('car_number','')},{r.get('freon_volume',0)},{r.get('money',0)},{r.get('payment_type','')},{r.get('payment_name','')},{r.get('payment_bank','')}\n")
            total_freon = sum(r.get("freon_volume", 0) for r in records)
            total_money = sum(r.get("money", 0) for r in records)
            f.write(f"ИТОГО,,,{len(records)} машин,{total_freon},{total_money},\n")
        return filename
